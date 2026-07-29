from __future__ import annotations

import hashlib
import html
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = ROOT / "work" / "blog-import" / "tu-lieu-bai-viet-website.docx"
MEDIA_DIR = ROOT / "frontend" / "public" / "media" / "blog"
OUTPUT_SQL = ROOT / "work" / "blog-import" / "import-blog-posts.sql"
GOOGLE_SOURCE = "https://docs.google.com/document/d/1qxTCi-hf3DLqMrN_qbIP7AbExDd4W-I24sh6PPXP32M"


@dataclass(frozen=True)
class Enrichment:
    category_id: int
    category: str
    deep_title: str
    deep_paragraphs: tuple[str, ...]
    tips_title: str
    tips: tuple[str, ...]
    faq: tuple[tuple[str, str], ...]
    tags: tuple[str, ...]


ENRICHMENTS: dict[int, Enrichment] = {
    1: Enrichment(
        17, "Giới thiệu", "Một điểm dừng mang tinh thần của thành phố biển",
        (
            "Một trải nghiệm ẩm thực đáng nhớ không chỉ nằm ở món ăn ngon mà còn ở cảm giác mà thương hiệu để lại. Với ALOO, ly kem bơ là điểm bắt đầu cho một khoảng nghỉ nhẹ nhàng giữa hành trình khám phá Quy Nhơn: đủ mát để xua đi cái nóng, đủ gần gũi để cả gia đình cùng thưởng thức và đủ đặc trưng để du khách nhớ về thành phố biển.",
            "ALOO phát triển trải nghiệm từ ba yếu tố có liên hệ chặt chẽ: nguyên liệu được lựa chọn kỹ, công thức cân bằng giữa bơ và kem dừa, cùng phong cách phục vụ thân thiện. Khi ba yếu tố này được duy trì đồng đều, khách hàng có thể dễ dàng nhận ra hương vị quen thuộc dù ghé cửa hàng vào những thời điểm khác nhau.",
            "Không gian trẻ trung và menu nhiều lựa chọn giúp ALOO phù hợp với nhiều nhịp ghé quán: một món tráng miệng sau bữa tối, điểm hẹn nhanh cùng bạn bè hoặc nơi nghỉ chân trong lịch trình du lịch. Chính sự linh hoạt ấy khiến kem bơ không còn là món ăn theo mùa mà trở thành một phần của trải nghiệm địa phương."
        ),
        "Gợi ý để lần đầu ghé ALOO thêm trọn vẹn",
        (
            "Bắt đầu với kem bơ truyền thống để cảm nhận rõ độ dẻo của bơ và vị mát của kem dừa.",
            "Nếu đi theo nhóm, gọi thêm một món vị đậm và một món thanh nhẹ để cùng chia sẻ.",
            "Thưởng thức ngay khi món vừa được phục vụ để giữ kết cấu mịn, lạnh và phần topping giòn.",
            "Kết hợp chuyến ghé quán với lịch dạo biển hoặc tham quan khu trung tâm để hành trình thuận tiện hơn."
        ),
        (
            ("ALOO chỉ có kem bơ không?", "Không. Kem bơ là sản phẩm đặc trưng, bên cạnh đó menu còn có nhiều món kem, đồ uống và món từ trái cây để phù hợp với sở thích khác nhau."),
            ("Ai phù hợp ghé ALOO?", "Không gian và mức giá hướng đến nhóm khách rộng, từ học sinh, sinh viên, gia đình đến khách du lịch đi một mình hoặc theo nhóm.")
        ),
        ("ALOO Kem Bơ", "kem bơ Quy Nhơn", "ẩm thực Quy Nhơn", "điểm đến Quy Nhơn")
    ),
    2: Enrichment(
        14, "Review", "Review một ly kem bơ nên bắt đầu từ đâu?",
        (
            "Để đánh giá kem bơ khách quan, trước hết nên quan sát màu sắc và kết cấu của phần bơ. Bơ chín vừa thường cho cảm giác mịn, dẻo, không tách nước và không để lại vị xơ rõ rệt. Khi trộn cùng kem dừa, hỗn hợp cần giữ được độ sánh thay vì trở nên quá lỏng.",
            "Tiếp theo là sự cân bằng hương vị. Một ly kem bơ dễ ăn không nên chỉ nổi bật vị ngọt; vị béo tự nhiên của bơ, mùi thơm dịu của dừa và độ lạnh của kem cần nâng đỡ lẫn nhau. Topping đóng vai trò tạo tương phản về kết cấu, nhưng không nên lấn át phần bơ vốn là nhân vật chính.",
            "Cuối cùng là trải nghiệm tổng thể: tốc độ phục vụ, cách trình bày, sự sạch sẽ và cảm giác thoải mái khi ngồi lại. Đây là những chi tiết khiến một món ngon trở thành địa chỉ đáng quay lại, đặc biệt với du khách chỉ có ít thời gian ở Quy Nhơn."
        ),
        "Checklist review dành cho người lần đầu",
        (
            "Thử một muỗng riêng phần bơ trước khi trộn để cảm nhận nguyên liệu nền.",
            "Trộn nhẹ từ dưới lên để kem dừa và bơ hòa quyện mà topping vẫn còn độ giòn.",
            "Đánh giá cả độ ngọt, độ béo và hậu vị thay vì chỉ dựa vào ấn tượng đầu tiên.",
            "Nếu đi nhiều người, thử hai đến ba món để có góc nhìn đầy đủ hơn về menu."
        ),
        (
            ("Kem bơ ngon có nhất thiết phải thật ngọt?", "Không. Độ ngọt vừa phải thường giúp hương bơ và dừa rõ hơn, đồng thời tạo cảm giác dễ ăn đến muỗng cuối."),
            ("Nên ăn riêng hay trộn đều?", "Bạn có thể thử riêng từng lớp trước, sau đó trộn nhẹ để cảm nhận sự thay đổi về hương vị và kết cấu.")
        ),
        ("review ALOO", "review kem bơ", "ăn gì Quy Nhơn", "kem bơ ngon")
    ),
    3: Enrichment(
        14, "Review", "Vì sao kem bơ trở thành món phải thử tại Quy Nhơn?",
        (
            "Trong danh sách món giải nhiệt của thành phố biển, kem bơ có lợi thế nhờ sự kết hợp giữa trái cây tươi và cảm giác mát lạnh. Món ăn vừa đủ quen để dễ tiếp cận, vừa có bản sắc riêng khi mỗi quán lựa chọn giống bơ, tỷ lệ kem và topping khác nhau.",
            "Một địa chỉ kem bơ đáng lưu lại thường đáp ứng ba tiêu chí: nguyên liệu ổn định, vị ngọt cân bằng và vị trí thuận tiện. Với du khách, khả năng ghé nhanh sau khi tham quan khu trung tâm cũng quan trọng không kém chất lượng món ăn.",
            "ALOO hướng trải nghiệm theo tinh thần gần gũi: món được phục vụ nhanh, menu dễ chọn và có nhiều mức vị để cả nhóm cùng thưởng thức. Đây là lý do một lần ghé quán có thể trở thành điểm nhấn nhỏ nhưng đáng nhớ trong hành trình Quy Nhơn."
        ),
        "Cách chọn địa chỉ kem bơ phù hợp",
        (
            "Ưu tiên nơi công khai địa chỉ, giờ hoạt động và thông tin liên hệ rõ ràng.",
            "Xem menu trước để chọn món phù hợp với trẻ nhỏ, người thích ít ngọt hoặc người mê vị đậm.",
            "Chọn cửa hàng gần cung đường đang tham quan để tiết kiệm thời gian di chuyển.",
            "Vào giờ đông, có thể chọn mang đi hoặc đặt giao tận nơi nếu lịch trình gấp."
        ),
        (
            ("Nên ăn kem bơ vào thời điểm nào?", "Kem bơ phù hợp như món tráng miệng hoặc bữa phụ trong ngày; nên dùng ngay sau khi phục vụ để giữ độ lạnh và kết cấu."),
            ("Đi nhóm đông có thuận tiện không?", "Có. Menu đa dạng giúp nhóm dễ gọi nhiều món và chia sẻ, nhưng nên chủ động chọn giờ ghé phù hợp vào mùa du lịch.")
        ),
        ("địa chỉ kem bơ", "kem bơ Quy Nhơn", "đặc sản Quy Nhơn", "ALOO")
    ),
    4: Enrichment(
        12, "Câu chuyện thương hiệu", "Từ nguyên liệu quen thuộc đến công thức có dấu ấn",
        (
            "Bí mật của một ly kem bơ hấp dẫn không nằm ở một thành phần đơn lẻ mà ở cách các thành phần được kết nối. Bơ cần đạt độ chín phù hợp, kem dừa phải đủ mịn và topping cần giữ được độ giòn. Chỉ một mắt xích thiếu ổn định cũng có thể làm trải nghiệm thay đổi.",
            "Quá trình chuẩn hóa giúp hương vị không phụ thuộc hoàn toàn vào cảm tính của người pha chế. Từ khâu nhận bơ, phân loại, bảo quản đến định lượng khi phục vụ đều cần tiêu chí rõ ràng. Đây là phần công việc ít xuất hiện trước khách hàng nhưng quyết định chất lượng của từng ly kem.",
            "ALOO lựa chọn phát triển sản phẩm dựa trên sự chân thật của nguyên liệu thay vì tạo cảm giác bằng hương liệu quá mạnh. Nhờ vậy, vị bơ vẫn là trung tâm, kem dừa đóng vai trò làm mát và topping bổ sung kết cấu thay vì che lấp hương vị chính."
        ),
        "Ba lớp trải nghiệm trong một ly kem bơ",
        (
            "Lớp bơ dẻo mịn tạo nền vị béo tự nhiên và màu sắc đặc trưng.",
            "Lớp kem dừa lạnh giúp món ăn thanh hơn, dễ thưởng thức trong thời tiết nóng.",
            "Lớp topping giòn tạo nhịp vui miệng và giúp mỗi muỗng có cảm giác khác nhau.",
            "Cách trộn nhẹ quyết định sự hòa quyện mà vẫn giữ được nét riêng của từng lớp."
        ),
        (
            ("Vì sao độ chín của bơ quan trọng?", "Bơ non dễ nhạt và xơ, trong khi bơ quá chín có thể mất màu và hương vị tươi. Độ chín vừa giúp phần bơ mịn và cân bằng hơn."),
            ("Công thức riêng có nghĩa là món luôn giống hệt nhau?", "Nguyên liệu tự nhiên luôn có dao động nhất định, nhưng quy trình và định lượng giúp trải nghiệm được duy trì ổn định nhất có thể.")
        ),
        ("bí mật kem bơ", "công thức ALOO", "bơ sáp Tây Nguyên", "câu chuyện sản phẩm")
    ),
    5: Enrichment(
        18, "Mẹo hay", "Hiểu quả bơ qua từng dấu hiệu nhỏ",
        (
            "Chọn bơ ngon là kỹ năng kết hợp giữa quan sát và cảm nhận. Màu vỏ có thể cung cấp gợi ý nhưng không phải tiêu chí duy nhất vì mỗi giống bơ có đặc điểm khác nhau. Cần xem thêm độ căng của vỏ, trạng thái cuống, trọng lượng và cảm giác khi cầm.",
            "Bơ chín tự nhiên thường mềm dần tương đối đều. Nếu một vùng quá mềm trong khi phần còn lại cứng, quả có thể đã bị va đập hoặc chín không đồng đều. Khi mua nhiều, nên chọn các quả ở những mức độ chín khác nhau để sử dụng lần lượt, tránh chín rộ cùng lúc.",
            "Sau khi bơ đạt độ chín mong muốn, việc chuyển sang ngăn mát có thể giúp làm chậm quá trình chín. Phần bơ đã cắt cần hạn chế tiếp xúc với không khí và nên dùng sớm để giữ màu, mùi và kết cấu."
        ),
        "Mẹo bảo quản bơ trong gia đình",
        (
            "Để bơ xanh ở nơi khô thoáng, tránh ánh nắng trực tiếp và không bọc kín khi chưa chín.",
            "Kiểm tra mỗi ngày bằng lực ấn rất nhẹ, không bóp mạnh làm dập phần thịt bên trong.",
            "Khi bơ vừa chín, chuyển vào ngăn mát và ưu tiên sử dụng trong thời gian ngắn.",
            "Với bơ đã cắt, che kín bề mặt và ghi nhớ rằng màu sắc có thể thay đổi khi tiếp xúc không khí."
        ),
        (
            ("Lắc nghe hạt có phải bơ đã chín?", "Đây chỉ là một dấu hiệu tham khảo và không đúng với mọi giống bơ. Nên kết hợp thêm độ mềm, cuống và tình trạng vỏ."),
            ("Có nên ép bơ chín thật nhanh?", "Nên ưu tiên quá trình chín tự nhiên và kiểm tra thường xuyên để giữ kết cấu tốt hơn.")
        ),
        ("cách chọn bơ", "mẹo chọn bơ ngon", "bơ sáp", "bảo quản bơ")
    ),
    6: Enrichment(
        15, "Vận hành", "Quy trình thủ công vẫn cần tính kỷ luật",
        (
            "Thủ công không đồng nghĩa với tùy ý. Một quy trình làm kem ổn định cần kiểm soát từ vệ sinh dụng cụ, nhiệt độ nguyên liệu, định lượng đến thời gian làm lạnh. Sự tỉ mỉ ở từng công đoạn giúp thành phẩm giữ kết cấu mịn và hạn chế chênh lệch giữa các mẻ.",
            "Nguyên liệu tươi cần được tiếp nhận và sử dụng theo vòng đời rõ ràng. Việc phân loại ngay từ đầu giúp đội ngũ lựa chọn đúng quả cho từng thời điểm chế biến, đồng thời giảm lãng phí do bơ chín quá mức hoặc chưa đạt yêu cầu.",
            "Khi đưa sản phẩm từ khu sản xuất đến cửa hàng, chuỗi bảo quản lạnh đóng vai trò quan trọng. Sự phối hợp giữa kho, vận chuyển và nhân viên cửa hàng giúp món được phục vụ ở trạng thái tốt, đặc biệt trong những ngày nắng nóng hoặc giờ cao điểm."
        ),
        "Những điểm kiểm soát chất lượng quan trọng",
        (
            "Vệ sinh bề mặt, dụng cụ và khu vực chế biến trước mỗi mẻ.",
            "Phân loại bơ theo độ chín và loại bỏ nguyên liệu không đạt yêu cầu cảm quan.",
            "Tuân thủ định lượng để vị ngọt, độ béo và kết cấu ổn định.",
            "Theo dõi nhiệt độ bảo quản và thời gian lưu trữ của từng nhóm nguyên liệu."
        ),
        (
            ("Kem thủ công có phải làm hoàn toàn bằng tay?", "Thủ công nhấn mạnh sự kiểm soát trực tiếp và công thức riêng; một số thiết bị vẫn được sử dụng để bảo đảm vệ sinh, độ mịn và nhiệt độ."),
            ("Không dùng chất bảo quản thì cần lưu ý gì?", "Quy trình phải chú trọng nguyên liệu tươi, nhiệt độ, thời gian lưu trữ và kế hoạch sản xuất phù hợp nhu cầu.")
        ),
        ("quy trình làm kem", "kem thủ công", "vận hành ALOO", "chất lượng sản phẩm")
    ),
    7: Enrichment(
        14, "Review", "Đọc menu theo khẩu vị thay vì chọn theo tên gọi",
        (
            "Một menu tốt giúp khách hàng nhận ra món phù hợp chỉ sau vài phút. Người thích vị nguyên bản có thể bắt đầu từ kem bơ truyền thống; người muốn hương vị mạnh hơn có thể chọn món kết hợp sầu riêng; còn nhóm đi cùng trẻ nhỏ thường ưu tiên các lựa chọn quen thuộc và dễ ăn.",
            "Sự khác biệt giữa các món không chỉ nằm ở topping. Tỷ lệ bơ, loại kem ăn kèm và kết cấu của thành phần phụ có thể làm trải nghiệm thay đổi rõ rệt. Vì vậy, gọi nhiều món để chia sẻ là cách thú vị để khám phá menu mà không bị quá no.",
            "Đồ uống và món trái cây đóng vai trò cân bằng cho nhóm khách có nhu cầu khác nhau. Khi một người muốn món béo, người khác muốn đồ uống thanh nhẹ, cả nhóm vẫn có thể ngồi cùng một không gian và tìm được lựa chọn phù hợp."
        ),
        "Gợi ý chọn món theo nhóm khách",
        (
            "Người lần đầu: ưu tiên món truyền thống để nhận biết hương vị cốt lõi của ALOO.",
            "Người thích vị đậm: chọn món kết hợp sầu riêng hoặc topping có độ béo rõ.",
            "Gia đình có trẻ nhỏ: hỏi nhân viên về độ ngọt và chọn khẩu phần phù hợp.",
            "Nhóm bạn: gọi nhiều món theo các nhóm vị khác nhau rồi cùng chia sẻ."
        ),
        (
            ("Món nào đại diện rõ nhất cho ALOO?", "Kem bơ truyền thống là lựa chọn phù hợp để cảm nhận nền bơ và kem dừa; món đặc biệt phù hợp khi muốn trải nghiệm nhiều lớp topping hơn."),
            ("Menu có phù hợp người không thích bơ?", "Bên cạnh kem bơ, ALOO còn có các lựa chọn kem và đồ uống khác tùy từng cửa hàng và thời điểm.")
        ),
        ("menu ALOO", "kem bơ đặc biệt", "review thực đơn", "món ngon Quy Nhơn")
    ),
    8: Enrichment(
        18, "Mẹo hay", "Thưởng thức quả bơ theo cách cân bằng",
        (
            "Quả bơ được yêu thích nhờ kết cấu béo mịn và khả năng kết hợp với nhiều món. Tuy vậy, cảm giác “tốt cho sức khỏe” không có nghĩa là cần ăn không giới hạn. Khẩu phần, tổng năng lượng trong ngày và các thành phần đi kèm vẫn là những yếu tố nên được cân nhắc.",
            "Khi dùng bơ trong món tráng miệng, độ ngọt thường đến từ kem, sữa hoặc topping chứ không chỉ từ quả bơ. Người quan tâm đến chế độ ăn có thể chọn khẩu phần vừa phải, thưởng thức chậm và xem món như một phần của tổng thể bữa ăn.",
            "Nội dung dinh dưỡng trên website chỉ mang tính tham khảo chung, không thay thế tư vấn cá nhân. Người có dị ứng, bệnh nền hoặc chế độ ăn đặc biệt nên trao đổi với chuyên gia y tế về lựa chọn phù hợp."
        ),
        "Cách đưa bơ vào thực đơn hợp lý",
        (
            "Chú ý khẩu phần và tần suất thay vì gắn nhãn một món là hoàn toàn tốt hoặc xấu.",
            "Kết hợp đa dạng rau, trái cây và nguồn đạm trong ngày thay vì chỉ tập trung vào một thực phẩm.",
            "Khi gọi món tráng miệng, có thể chia sẻ cùng người thân để vừa trải nghiệm vừa cân bằng khẩu phần.",
            "Lắng nghe cảm giác no và nhu cầu thực tế của cơ thể."
        ),
        (
            ("Ăn bơ có chắc chắn giúp giảm cân?", "Không có một thực phẩm đơn lẻ quyết định việc giảm cân. Kết quả phụ thuộc tổng năng lượng, vận động, giấc ngủ và nhiều yếu tố cá nhân."),
            ("Kem bơ có giống ăn bơ tươi?", "Không hoàn toàn. Kem bơ còn có kem, sữa hoặc topping nên thành phần và năng lượng khác với bơ tươi.")
        ),
        ("lợi ích quả bơ", "ăn bơ cân bằng", "kem bơ", "mẹo dinh dưỡng")
    ),
    9: Enrichment(
        12, "Câu chuyện thương hiệu", "Một thương hiệu lớn lên từ giá trị nhỏ được làm đều đặn",
        (
            "Hành trình thương hiệu thường không bắt đầu bằng quy mô lớn mà bằng một sản phẩm được khách hàng tin tưởng. Với ALOO, nền tảng ấy là ly kem bơ từ nguyên liệu gần gũi, được hoàn thiện qua phản hồi thực tế của người địa phương và du khách.",
            "Khi số lượng cửa hàng tăng, thách thức không chỉ là bán nhiều hơn mà còn là giữ được tinh thần ban đầu. Quy trình, đào tạo và tiêu chuẩn phục vụ trở thành cách để chuyển những kinh nghiệm cá nhân thành năng lực của cả hệ thống.",
            "Tầm nhìn phát triển của ALOO gắn với việc nâng giá trị nông sản Việt thông qua sản phẩm dễ tiếp cận. Mỗi món mới cần vừa tạo hứng thú cho khách hàng, vừa duy trì mối liên hệ với nguyên liệu và câu chuyện mà thương hiệu theo đuổi."
        ),
        "Những cột mốc tạo nên bản sắc ALOO",
        (
            "Khởi đầu từ nhu cầu mang đến món giải nhiệt gần gũi cho cộng đồng địa phương.",
            "Hoàn thiện công thức dựa trên phản hồi và thói quen thưởng thức thực tế.",
            "Chuẩn hóa vận hành để hương vị và dịch vụ có thể được nhân rộng.",
            "Mở rộng hoạt động cộng đồng để thương hiệu tạo ra giá trị ngoài sản phẩm."
        ),
        (
            ("Điều gì tạo nên bản sắc ALOO?", "Sự kết hợp giữa nguyên liệu bơ, công thức kem dừa, trải nghiệm thân thiện và câu chuyện phát triển từ địa phương."),
            ("Vì sao thương hiệu cần chuẩn hóa?", "Chuẩn hóa giúp kiến thức không phụ thuộc một cá nhân và tạo nền tảng để nhiều cửa hàng phục vụ ổn định hơn.")
        ),
        ("câu chuyện ALOO", "thương hiệu địa phương", "nông sản Việt", "hành trình thương hiệu")
    ),
    10: Enrichment(
        12, "Câu chuyện thương hiệu", "Những khoảnh khắc nhỏ tạo nên ký ức về quán",
        (
            "Một cửa hàng đồ ăn không chỉ lưu giữ hình ảnh món ăn mà còn có những cuộc gặp gỡ: nhóm bạn hẹn nhau sau giờ học, gia đình dừng chân sau chuyến đi biển, hay du khách hỏi nhân viên về một địa điểm nên ghé tiếp theo. Những khoảnh khắc đời thường ấy tạo nên sức sống cho không gian.",
            "Hình ảnh khách hàng có giá trị khi thể hiện sự tự nhiên và tôn trọng quyền riêng tư. ALOO ưu tiên những khung hình ghi lại niềm vui, sự kết nối và tinh thần tích cực, đồng thời cần có sự đồng thuận phù hợp trước khi sử dụng cho truyền thông.",
            "Qua thời gian, bộ ảnh tại quán trở thành một cuốn nhật ký cộng đồng. Mỗi mùa du lịch, mỗi hoạt động nhỏ và mỗi nụ cười đều bổ sung thêm một lớp ký ức cho câu chuyện thương hiệu."
        ),
        "Gợi ý lưu lại khoảnh khắc đẹp tại ALOO",
        (
            "Chụp món ngay khi vừa phục vụ để giữ màu sắc và kết cấu đẹp nhất.",
            "Ưu tiên ánh sáng tự nhiên và khung hình gọn để món ăn trở thành điểm nhấn.",
            "Nếu chụp cùng nhân viên hoặc khách khác, hãy xin phép trước khi đăng tải.",
            "Gắn vị trí cửa hàng và chia sẻ cảm nhận thật để bài đăng hữu ích hơn cho cộng đồng."
        ),
        (
            ("ALOO có thể chia sẻ ảnh của khách hàng không?", "Việc sử dụng ảnh cho truyền thông cần phù hợp với quyền riêng tư và sự đồng thuận của người xuất hiện trong ảnh."),
            ("Làm sao để ảnh món ăn trông tự nhiên?", "Dùng ánh sáng mềm, hạn chế bộ lọc quá mạnh và chụp ngay khi món vừa được phục vụ.")
        ),
        ("khoảnh khắc khách hàng", "cộng đồng ALOO", "check-in Quy Nhơn", "câu chuyện tại quán")
    ),
    11: Enrichment(
        14, "Review", "Vì sao góc nhìn người bản địa đáng tham khảo?",
        (
            "Người địa phương thường đánh giá một quán qua sự ổn định lâu dài thay vì ấn tượng của một lần ghé. Họ biết thời điểm quán đông, món nào dễ ăn, cách kết hợp phù hợp và liệu chất lượng có được duy trì qua nhiều mùa hay không.",
            "Một review hữu ích nên nói rõ bối cảnh: đi cùng ai, gọi món gì, thời điểm nào và khẩu vị cá nhân ra sao. Những chi tiết này giúp người đọc hiểu trải nghiệm có phù hợp với mình thay vì xem nhận xét như một kết luận tuyệt đối.",
            "ALOO trân trọng phản hồi tích cực lẫn góp ý cụ thể. Review không chỉ hỗ trợ khách mới lựa chọn mà còn là nguồn dữ liệu để cửa hàng điều chỉnh tốc độ phục vụ, cách giới thiệu món và chất lượng trải nghiệm."
        ),
        "Cách đọc review thông minh",
        (
            "Ưu tiên nhận xét có mô tả món, thời điểm và trải nghiệm cụ thể.",
            "So sánh nhiều góc nhìn thay vì dựa vào một đánh giá quá tích cực hoặc quá tiêu cực.",
            "Phân biệt sở thích cá nhân, chẳng hạn thích rất ngọt, với vấn đề chất lượng thực tế.",
            "Kiểm tra thông tin địa chỉ và menu mới nhất trước khi di chuyển."
        ),
        (
            ("Review của người bản địa có luôn chính xác?", "Đây là nguồn tham khảo giàu trải nghiệm, nhưng khẩu vị vẫn mang tính cá nhân và thông tin có thể thay đổi theo thời điểm."),
            ("ALOO tiếp nhận góp ý ở đâu?", "Khách hàng có thể phản hồi trực tiếp tại cửa hàng hoặc qua các kênh liên hệ chính thức của thương hiệu.")
        ),
        ("review người bản địa", "ALOO Quy Nhơn", "kinh nghiệm ăn uống", "kem bơ địa phương")
    ),
    12: Enrichment(
        17, "Giới thiệu", "Một điểm hẹn tốt cần điều gì?",
        (
            "Với nhóm bạn, địa điểm tụ tập lý tưởng cần dung hòa nhiều nhu cầu: dễ tìm, menu có nhiều mức vị, thời gian phục vụ hợp lý và không khí đủ thoải mái để trò chuyện. Một món ăn ngon là lý do để đến, nhưng cảm giác thuận tiện mới là lý do để cả nhóm quay lại.",
            "Kem bơ phù hợp để chia sẻ vì khẩu phần gọn và có nhiều biến thể. Mỗi người có thể chọn một món riêng hoặc gọi theo nhóm rồi đổi vị cho nhau. Cách thưởng thức này tạo thêm tương tác và giúp cuộc gặp trở nên vui hơn.",
            "ALOO hướng đến không gian trẻ trung, gần gũi với cả nhóm học sinh, sinh viên, gia đình và khách du lịch. Dù ghé nhanh hay ngồi lại, khách hàng vẫn có thể chủ động chọn cách trải nghiệm phù hợp với lịch trình."
        ),
        "Lên kèo nhóm bạn thật gọn",
        (
            "Chọn cửa hàng thuận đường cho phần lớn thành viên và gửi vị trí trước khi hẹn.",
            "Thống nhất khung giờ, đặc biệt vào cuối tuần hoặc mùa cao điểm du lịch.",
            "Gọi nhiều nhóm vị để mọi người có cơ hội thử món mới.",
            "Nếu tổ chức sinh nhật nhỏ, nên liên hệ cửa hàng trước để hỏi khả năng sắp xếp."
        ),
        (
            ("Nhóm đông có cần đặt trước không?", "Với nhóm đông hoặc dịp đặc biệt, liên hệ trước giúp cửa hàng tư vấn thời điểm và cách sắp xếp phù hợp."),
            ("Có lựa chọn mang đi không?", "ALOO phục vụ nhiều hình thức tùy cửa hàng, gồm dùng tại chỗ, mang đi và giao hàng qua kênh phù hợp.")
        ),
        ("tụ tập nhóm bạn", "điểm hẹn Quy Nhơn", "ALOO Kem Bơ", "ăn vặt nhóm")
    ),
    13: Enrichment(
        19, "Khuyến mãi", "Ưu đãi chỉ thật sự tốt khi thông tin minh bạch",
        (
            "Một chương trình cuối tuần hấp dẫn cần giúp khách hàng hiểu ngay ba điều: được ưu đãi gì, áp dụng khi nào và có điều kiện nào đi kèm. Thông tin rõ ràng giúp khách chủ động lựa chọn, đồng thời giảm hiểu nhầm tại quầy.",
            "Ưu đãi nên khuyến khích khách khám phá menu hoặc chia sẻ trải nghiệm cùng người thân thay vì chỉ tập trung vào giảm giá. Các combo nhóm, quà tặng nhỏ hoặc gợi ý thử món mới có thể tạo giá trị thiết thực hơn cho chuyến ghé quán.",
            "Do chương trình có thể thay đổi theo từng cửa hàng và thời điểm, khách hàng nên kiểm tra kênh chính thức trước khi đến. ALOO không khuyến khích dựa vào ảnh chụp cũ hoặc thông tin truyền miệng chưa được xác nhận."
        ),
        "Cách kiểm tra ưu đãi cuối tuần",
        (
            "Xem ngày bắt đầu, ngày kết thúc và khung giờ áp dụng.",
            "Kiểm tra cửa hàng tham gia và hình thức dùng tại chỗ, mang đi hoặc giao hàng.",
            "Đọc điều kiện về số lượng, hóa đơn tối thiểu hoặc sản phẩm áp dụng.",
            "Lưu bài đăng chính thức để đối chiếu khi cần."
        ),
        (
            ("Ưu đãi có áp dụng ở mọi cửa hàng không?", "Không nhất thiết. Phạm vi áp dụng phụ thuộc nội dung công bố của từng chương trình."),
            ("Có cộng dồn nhiều ưu đãi không?", "Khả năng cộng dồn cần được nêu trong thể lệ; nếu chưa rõ, khách hàng nên hỏi cửa hàng trước khi thanh toán.")
        ),
        ("ưu đãi ALOO", "khuyến mãi cuối tuần", "combo kem bơ", "tin ALOO")
    ),
    14: Enrichment(
        20, "Sự kiện cộng đồng", "Mini game hay bắt đầu từ một thể lệ dễ hiểu",
        (
            "Mini game là cách tạo không khí vui vẻ và khuyến khích cộng đồng tương tác, nhưng trải nghiệm chỉ tích cực khi thể lệ ngắn gọn, công bằng và có thời hạn rõ ràng. Người tham gia cần biết chính xác hành động phải thực hiện và cách kết quả được xác định.",
            "Nội dung thử thách nên gắn với câu chuyện thương hiệu một cách tự nhiên, chẳng hạn chia sẻ khoảnh khắc cùng bạn bè, kể lại món yêu thích hoặc sáng tạo lời nhắn tích cực. Khi người tham gia có không gian thể hiện cá tính, mini game trở thành hoạt động cộng đồng thay vì một yêu cầu chia sẻ máy móc.",
            "ALOO ưu tiên công bố kết quả trên kênh chính thức và liên hệ người nhận quà bằng phương thức minh bạch. Khách hàng cần cảnh giác với tài khoản giả mạo yêu cầu cung cấp mật khẩu, mã xác thực hoặc chuyển tiền để nhận thưởng."
        ),
        "Checklist tham gia mini game an toàn",
        (
            "Đọc kỹ thể lệ, thời hạn và điều kiện hợp lệ trước khi tham gia.",
            "Chỉ tương tác với trang và kênh liên hệ chính thức của ALOO.",
            "Không cung cấp mật khẩu, mã OTP hoặc chuyển phí nhận giải.",
            "Kiểm tra bài công bố kết quả và hướng dẫn nhận quà chính thức."
        ),
        (
            ("ALOO liên hệ người trúng giải bằng cách nào?", "Phương thức liên hệ phải được nêu trong thể lệ hoặc bài công bố chính thức của chương trình."),
            ("Có phải trả phí để nhận quà không?", "Không nên chuyển tiền cho tài khoản không xác minh. Hãy đối chiếu thông tin trực tiếp với kênh chính thức của ALOO.")
        ),
        ("mini game ALOO", "sự kiện ALOO", "quà tặng", "cộng đồng ALOO")
    ),
    15: Enrichment(
        18, "Mẹo hay", "Chọn món giải nhiệt cho trẻ nhỏ bằng sự cân bằng",
        (
            "Khi chọn món cho trẻ, phụ huynh thường quan tâm đến nguyên liệu, khẩu phần, độ ngọt và khả năng dị ứng. Một món có thành phần quen thuộc vẫn cần được lựa chọn theo độ tuổi, thói quen ăn uống và tình trạng sức khỏe riêng của từng bé.",
            "Kem bơ có kết cấu mềm và hương vị dễ tiếp cận, nhưng phần kem, sữa và topping có thể làm tăng độ ngọt hoặc chứa thành phần mà trẻ nhạy cảm. Phụ huynh nên hỏi nhân viên về nguyên liệu và cân nhắc khẩu phần nhỏ khi bé thử lần đầu.",
            "Website chỉ cung cấp thông tin trải nghiệm chung, không thay thế tư vấn y tế. Trẻ có tiền sử dị ứng thực phẩm, không dung nạp sữa hoặc cần chế độ ăn đặc biệt nên được người chăm sóc kiểm tra kỹ trước khi dùng."
        ),
        "Gợi ý để bé thưởng thức vui và an toàn hơn",
        (
            "Chọn khẩu phần phù hợp và cho bé ăn chậm thay vì dùng quá nhanh khi món còn rất lạnh.",
            "Hỏi về thành phần sữa, dừa, đậu phộng và các loại hạt trong topping.",
            "Không ép trẻ ăn hết; theo dõi phản ứng khi thử món hoặc nguyên liệu mới.",
            "Kết hợp buổi ăn kem với nước lọc và lịch nghỉ ngơi phù hợp trong ngày du lịch."
        ),
        (
            ("Kem bơ có phù hợp với mọi trẻ nhỏ?", "Không thể khẳng định cho mọi trẻ. Phụ huynh cần cân nhắc độ tuổi, khẩu phần, dị ứng và hướng dẫn dinh dưỡng cá nhân."),
            ("Có thể bỏ topping không?", "Tùy món và cửa hàng, phụ huynh có thể trao đổi với nhân viên để chọn cách phục vụ phù hợp hơn.")
        ),
        ("kem bơ cho bé", "gia đình du lịch", "món giải nhiệt", "ALOO gia đình")
    ),
    16: Enrichment(
        18, "Mẹo hay", "Kem bơ có béo không: nhìn vào cả món, không chỉ quả bơ",
        (
            "Câu hỏi “có béo không” thường quá đơn giản so với thực tế. Thay đổi cân nặng phụ thuộc tổng năng lượng trong thời gian dài, mức vận động, giấc ngủ và nhiều yếu tố cá nhân. Một phần kem bơ không tự động quyết định kết quả nếu được đặt trong chế độ ăn tổng thể.",
            "Bơ có vị béo tự nhiên, trong khi kem, sữa và topping bổ sung thêm năng lượng. Vì vậy, kem bơ là món tráng miệng khác với việc ăn bơ tươi. Người muốn kiểm soát khẩu phần có thể chọn phần vừa, chia sẻ cùng bạn bè và chú ý các món khác trong ngày.",
            "Không cần xem món tráng miệng như điều phải né tránh hoàn toàn. Thưởng thức có chủ đích, ăn chậm và chọn tần suất phù hợp thường bền vững hơn tư duy “ăn bù” hoặc kiêng tuyệt đối."
        ),
        "Bốn cách thưởng thức chủ động hơn",
        (
            "Chọn khẩu phần phù hợp với mức đói và nhu cầu của bản thân.",
            "Thưởng thức chậm để cảm nhận hương vị thay vì ăn theo quán tính.",
            "Chia sẻ topping hoặc món đặc biệt khi muốn thử nhiều vị.",
            "Không sử dụng bài viết này thay cho tư vấn dinh dưỡng hoặc điều trị cá nhân."
        ),
        (
            ("Ăn kem bơ buổi tối có chắc chắn tăng cân?", "Thời điểm không phải yếu tố duy nhất. Tổng năng lượng và thói quen trong thời gian dài quan trọng hơn một lần ăn."),
            ("Có phiên bản ít ngọt không?", "Khách hàng có thể hỏi nhân viên về lựa chọn và cách phục vụ tại từng cửa hàng, nhưng công thức cụ thể phụ thuộc món.")
        ),
        ("kem bơ có béo không", "ăn kem cân bằng", "khẩu phần kem bơ", "mẹo dinh dưỡng")
    ),
    17: Enrichment(
        14, "Review", "Thực đơn “bí mật” là cách gọi cho trải nghiệm khám phá",
        (
            "Sự hấp dẫn của một món ít người biết đến nằm ở cảm giác khám phá, nhưng thông tin vẫn cần rõ ràng về nguyên liệu và giá bán. “Bí mật” nên được hiểu là gợi ý kết hợp thú vị hoặc món nổi bật theo thời điểm, không phải sản phẩm thiếu thông tin.",
            "Những món mix thường tạo khác biệt bằng cách ghép hai nhóm hương vị: vị béo mịn của bơ với sầu riêng đậm, hoặc nền bơ với kem dừa và topping giòn. Tỷ lệ hợp lý giúp món có chiều sâu nhưng không khiến người ăn nhanh ngấy.",
            "Khách hàng nên hỏi nhân viên về tình trạng món và thành phần trước khi gọi, nhất là khi có dị ứng với sữa, dừa, đậu phộng hoặc các loại hạt. Menu có thể thay đổi theo mùa và khả năng cung ứng nguyên liệu."
        ),
        "Cách khám phá món mix không bị quá tải vị",
        (
            "Bắt đầu từ món có một hương vị mới thay vì chọn quá nhiều topping cùng lúc.",
            "Đi theo nhóm để chia sẻ và so sánh nhiều món mà vẫn giữ khẩu phần vừa phải.",
            "Dùng một ít nước lọc giữa các món để cảm nhận rõ sự khác nhau.",
            "Hỏi nhân viên về mức độ ngọt và độ đậm của món trước khi lựa chọn."
        ),
        (
            ("Món “bí mật” có luôn sẵn không?", "Không nhất thiết. Một số món hoặc cách kết hợp có thể phụ thuộc nguyên liệu, mùa và cửa hàng."),
            ("Có thể tự chọn topping không?", "Khả năng tùy chỉnh tùy món và chính sách phục vụ tại từng cửa hàng; khách hàng nên hỏi trực tiếp.")
        ),
        ("thực đơn bí mật", "kem bơ mix", "review menu ALOO", "món đặc biệt")
    ),
    18: Enrichment(
        20, "Sự kiện cộng đồng", "Khi một ly kem trở thành nhịp cầu sẻ chia",
        (
            "Hoạt động cộng đồng có ý nghĩa khi tạo ra sự kết nối thật giữa doanh nghiệp, người tham gia và địa phương. Với ALOO, việc mang sản phẩm đến gần trẻ em không chỉ là trao một món quà mà còn là dịp để lắng nghe, gặp gỡ và chia sẻ niềm vui.",
            "Mỗi chương trình cần được chuẩn bị có trách nhiệm: phối hợp với đơn vị địa phương, bảo đảm an toàn thực phẩm, lựa chọn cách trao tặng phù hợp và tôn trọng hình ảnh của trẻ em. Truyền thông nên tập trung vào giá trị chung thay vì biến người nhận thành công cụ quảng bá.",
            "Dấu ấn bền vững không đến từ một sự kiện đơn lẻ mà từ cam kết đồng hành lâu dài. Những trải nghiệm tại Khánh Sơn giúp đội ngũ ALOO hiểu hơn về nhu cầu cộng đồng và cách thương hiệu có thể đóng góp bằng nguồn lực thực tế."
        ),
        "Nguyên tắc của một hoạt động cộng đồng tử tế",
        (
            "Phối hợp với tổ chức hoặc đầu mối địa phương để hiểu đúng nhu cầu.",
            "Bảo đảm chất lượng sản phẩm và quy trình vận chuyển trong toàn bộ chương trình.",
            "Tôn trọng quyền riêng tư, đặc biệt khi chụp và sử dụng hình ảnh trẻ em.",
            "Công khai thông tin vừa đủ, tránh phóng đại tác động hoặc biến sẻ chia thành khẩu hiệu."
        ),
        (
            ("Hoạt động này có diễn ra thường xuyên không?", "Thông tin từng chương trình sẽ được ALOO cập nhật trên các kênh chính thức khi có kế hoạch cụ thể."),
            ("Làm sao đồng hành cùng chương trình?", "Cá nhân hoặc tổ chức quan tâm có thể liên hệ ALOO qua kênh chính thức để trao đổi hình thức phù hợp.")
        ),
        ("ALOO Nha Trang", "Khánh Sơn", "hoạt động cộng đồng", "yêu thương")
    ),
    19: Enrichment(
        20, "Sự kiện cộng đồng", "Một gian hàng sự kiện cần nhiều hơn sản phẩm ngon",
        (
            "Tại sự kiện đông người, trải nghiệm khách hàng phụ thuộc vào khả năng chuẩn bị và phối hợp. Sản phẩm cần được bảo quản đúng điều kiện, khu vực phục vụ phải gọn sạch và đội ngũ cần xử lý hàng chờ một cách thân thiện, rõ ràng.",
            "Tíu Tít Paperfest tạo không gian để gia đình và người trẻ cùng sáng tạo. Sự hiện diện của ALOO bổ sung một điểm nghỉ mát lạnh giữa các hoạt động, nơi khách có thể nạp năng lượng, trò chuyện và tiếp tục hành trình khám phá.",
            "Đối với thương hiệu, sự kiện là dịp thử nghiệm năng lực vận hành ngoài cửa hàng. Những phản hồi về món, tốc độ phục vụ và cách bố trí gian hàng trở thành dữ liệu hữu ích để ALOO hoàn thiện cho các chương trình tiếp theo."
        ),
        "Kinh nghiệm tham gia Paperfest cùng gia đình",
        (
            "Xem trước thời gian và địa điểm chính thức để chủ động lộ trình.",
            "Chuẩn bị nước uống, đồ chống nắng và khoảng nghỉ phù hợp cho trẻ nhỏ.",
            "Ưu tiên trải nghiệm hoạt động trước, sau đó nghỉ và dùng món ở thời điểm thoải mái.",
            "Giữ vệ sinh chung và phân loại rác theo hướng dẫn của ban tổ chức."
        ),
        (
            ("ALOO mang món gì đến sự kiện?", "Danh mục phục vụ tùy kế hoạch chương trình và điều kiện vận hành; bài viết nguồn giới thiệu các món kem bơ và lựa chọn giải nhiệt đặc trưng."),
            ("Có cần mua vé để ghé gian hàng không?", "Khách nên kiểm tra thông tin chính thức của ban tổ chức về quyền vào cửa và hoạt động áp dụng.")
        ),
        ("Tíu Tít Paperfest 2026", "sự kiện Quy Nhơn", "gian hàng ALOO", "năng lượng xanh")
    ),
    20: Enrichment(
        20, "Sự kiện cộng đồng", "Sáng tạo và năng lượng xanh gặp nhau như thế nào?",
        (
            "Sáng tạo thường bắt đầu từ những vật liệu và ý tưởng rất gần gũi. Paperfest khuyến khích người tham gia nhìn một tờ giấy bằng góc nhìn mới; ALOO cũng theo đuổi tinh thần tương tự khi biến quả bơ quen thuộc thành nhiều trải nghiệm ẩm thực khác nhau.",
            "Điểm chung giữa hai hành trình là sự dám thử và khả năng hoàn thiện qua từng lần thực hiện. Một sản phẩm đẹp hay một món ngon đều cần ý tưởng, kỹ năng, kỷ luật và sự kiên nhẫn. Vì thế, gian hàng ALOO không chỉ là nơi giải nhiệt mà còn góp phần nối dài câu chuyện sáng tạo của ngày hội.",
            "Năng lượng xanh trong bối cảnh này được hiểu theo nghĩa tích cực và gần gũi: trân trọng nguyên liệu, ưu tiên trải nghiệm cộng đồng, giữ không gian sạch và khuyến khích lối sống chủ động. Đây là thông điệp ALOO muốn duy trì sau khi sự kiện kết thúc."
        ),
        "Mang tinh thần Paperfest về cuộc sống hằng ngày",
        (
            "Bắt đầu một ý tưởng nhỏ thay vì chờ điều kiện hoàn hảo.",
            "Tận dụng vật liệu hợp lý và dọn sạch không gian sau mỗi hoạt động.",
            "Chia sẻ thành quả, góp ý tử tế và ghi nhận công sức của người khác.",
            "Dành thời gian nghỉ, nạp năng lượng và quay lại với ý tưởng bằng tâm thế tươi mới."
        ),
        (
            ("Sự kiện để lại điều gì cho ALOO?", "Đây là cơ hội kết nối cộng đồng, kiểm chứng năng lực phục vụ ngoài cửa hàng và tiếp nhận phản hồi trực tiếp."),
            ("Thông tin sự kiện được cập nhật ở đâu?", "Khách hàng nên theo dõi kênh chính thức của ALOO và ban tổ chức để nhận lịch trình, thể lệ và thay đổi mới nhất.")
        ),
        ("ALOO x Paperfest", "Paperfest Vietnam 2026", "sáng tạo", "năng lượng xanh")
    ),
}


def normalize_brand(text: str) -> str:
    replacements = {
        "ALoo": "ALOO",
        "Aloo": "ALOO",
        "Daklak": "Đắk Lắk",
        "Quy nhơn": "Quy Nhơn",
        "khoảng khắc": "khoảnh khắc",
        "Câu chuyệ ": "Câu chuyện ",
        "e - commerce": "thương mại điện tử",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: str) -> str:
    value = unicodedata.normalize("NFD", value.lower())
    value = "".join(character for character in value if unicodedata.category(character) != "Mn")
    value = value.replace("đ", "d")
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value[:270]


def paragraph_images(document: Document, paragraph) -> list[tuple[bytes, str]]:
    found: list[tuple[bytes, str]] = []
    for run in paragraph.runs:
        for blip in run._element.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed"))
            if not relationship_id:
                continue
            image_part = document.part.related_parts[relationship_id]
            suffix = Path(str(image_part.partname)).suffix.lower() or ".png"
            found.append((image_part.blob, suffix))
    return found


def is_heading(paragraph) -> bool:
    text = paragraph.text.strip()
    if not text or len(text) > 180:
        return False
    if re.match(r"^\d+[\.\)]\s+", text):
        return True
    if text.lower() in {"kết luận", "lời kết", "tổng kết"}:
        return True
    visible_runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(visible_runs) and all(run.bold for run in visible_runs)


def sql_literal(value: str | None, tag: str) -> str:
    if value is None:
        return "NULL"
    safe_tag = re.sub(r"[^a-zA-Z0-9_]", "_", tag)
    delimiter = f"${safe_tag}$"
    while delimiter in value:
        safe_tag += "_x"
        delimiter = f"${safe_tag}$"
    return f"{delimiter}{value}{delimiter}"


def article_title(raw: str) -> str:
    title = re.sub(r"^Bài\s*viết(?:\s*số)?\s*\d+\s*:\s*", "", raw, flags=re.IGNORECASE).strip()
    title = normalize_brand(title)
    title = re.sub(r":\s*(Nhấn mạnh|Tạo sự tò mò)\b[\s\S]*$", "", title, flags=re.IGNORECASE).strip()
    return title


def render_source(document: Document, paragraphs, article_number: int) -> tuple[str, list[str], str]:
    parts: list[str] = []
    image_urls: list[str] = []
    first_body = ""
    seen_hashes: set[str] = set()
    in_list = False

    for paragraph in paragraphs:
        raw = normalize_brand(paragraph.text)
        if raw.upper().startswith("MỌI CHI TIẾT XIN VUI LÒNG LIÊN HỆ"):
            break

        images = paragraph_images(document, paragraph)
        if raw:
            if not first_body and not is_heading(paragraph):
                first_body = raw
            list_item = re.match(r"^(?:[-•]\s*)(.+)$", raw)
            if list_item:
                if not in_list:
                    parts.append("<ul>")
                    in_list = True
                parts.append(f"<li>{html.escape(list_item.group(1))}</li>")
            else:
                if in_list:
                    parts.append("</ul>")
                    in_list = False
                if is_heading(paragraph):
                    heading = re.sub(r"^\d+[\.\)]\s*", "", raw).strip()
                    parts.append(f"<h2>{html.escape(heading)}</h2>")
                else:
                    parts.append(f"<p>{html.escape(raw)}</p>")

        for blob, suffix in images:
            digest = hashlib.sha256(blob).hexdigest()
            if digest in seen_hashes:
                continue
            seen_hashes.add(digest)
            image_index = len(image_urls) + 1
            filename = f"aloo-blog-{article_number:02d}-{image_index:02d}{suffix}"
            destination = MEDIA_DIR / filename
            destination.write_bytes(blob)
            image_url = f"/media/blog/{filename}"
            image_urls.append(image_url)
            alt = f"Hình ảnh bài viết {article_number} của ALOO"
            parts.append(
                f'<figure><img src="{image_url}" alt="{html.escape(alt)}" '
                'loading="lazy" decoding="async"></figure>'
            )

    if in_list:
        parts.append("</ul>")
    return "\n".join(parts), image_urls, first_body


def enrichment_html(item: Enrichment, title: str) -> str:
    parts = [f"<h2>{html.escape(item.deep_title)}</h2>"]
    parts.extend(f"<p>{html.escape(paragraph)}</p>" for paragraph in item.deep_paragraphs)
    parts.append(f"<h2>{html.escape(item.tips_title)}</h2>")
    parts.append("<ul>")
    parts.extend(f"<li>{html.escape(tip)}</li>" for tip in item.tips)
    parts.append("</ul>")
    parts.append("<h2>Câu hỏi thường gặp</h2>")
    for question, answer in item.faq:
        parts.append(f"<h3>{html.escape(question)}</h3>")
        parts.append(f"<p>{html.escape(answer)}</p>")
    parts.append("<h2>Gợi ý tiếp theo từ ALOO</h2>")
    parts.append(
        "<p>Sau khi tìm hiểu chủ đề này, bạn có thể khám phá thêm menu, câu chuyện nguyên liệu "
        "và hệ thống cửa hàng trên website ALOO. Thông tin sản phẩm, chương trình và lịch hoạt động "
        "có thể thay đổi theo thời điểm; hãy kiểm tra kênh chính thức trước khi lên lịch ghé quán.</p>"
    )
    parts.append(
        f"<p><strong>{html.escape(title)}</strong> không chỉ là một chủ đề để đọc nhanh, "
        "mà còn là lời mời trải nghiệm ALOO bằng sự tò mò, chậm rãi và cởi mở với những hương vị mới.</p>"
    )
    return "\n".join(parts)


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(f"Missing source document: {SOURCE_DOCX}")

    MEDIA_DIR.mkdir(parents=True, exist_ok=True)
    for old_image in MEDIA_DIR.glob("aloo-blog-*"):
        old_image.unlink()

    document = Document(SOURCE_DOCX)
    start_indexes = [
        index for index, paragraph in enumerate(document.paragraphs)
        if re.match(r"^Bài\s+(?:viết\s+)?(?:số\s+)?\d+\s*[-:]", paragraph.text.strip(), re.IGNORECASE)
        and (
            index == 0
            or not re.match(
                r"^Bài\s*viết(?:\s*số)?\s*\d+\s*:",
                document.paragraphs[index - 1].text.strip(),
                re.IGNORECASE,
            )
        )
    ]
    # Exported Google Docs tabs start with a tab title followed by the article title.
    start_indexes = [0, 39, 80, 123, 154, 189, 225, 259, 291, 326, 362, 398, 431, 464, 500, 533, 563, 595, 627, 658]
    if len(start_indexes) != 20:
        raise RuntimeError(f"Expected 20 articles, found {len(start_indexes)}")

    sql_parts = ["BEGIN;"]
    manifest_rows = ["number\ttitle\tslug\tcategory\timages\tcontent_chars"]

    for offset, start_index in enumerate(start_indexes):
        article_number = offset + 1
        end_index = start_indexes[offset + 1] if offset + 1 < len(start_indexes) else len(document.paragraphs)
        title_index = start_index + 1
        title = article_title(document.paragraphs[title_index].text)
        if article_number == 10:
            title = "Những khoảnh khắc khách hàng tại quán ALOO"
        source_html, image_urls, first_body = render_source(
            document,
            document.paragraphs[title_index + 1:end_index],
            article_number,
        )
        enrichment = ENRICHMENTS[article_number]
        content = f"{source_html}\n{enrichment_html(enrichment, title)}"
        excerpt = normalize_brand(first_body)[:420]
        if len(first_body) > 420:
            excerpt = excerpt.rsplit(" ", 1)[0] + "…"
        slug = slugify(title)
        thumbnail = image_urls[0] if image_urls else "/logo-aloo.png"
        published_at = f"2026-07-{article_number:02d} 08:00:00"
        seo_description = excerpt[:500]
        tags = ",".join(enrichment.tags)

        tag = f"blog_{article_number:02d}"
        sql_parts.append(
            f"""
INSERT INTO posts (
    title, slug, excerpt, content, thumbnail_url, category_id, author, source, source_link,
    article_type, status, published_at, seo_title, seo_description, meta_keywords,
    canonical_url, tags, created_at, updated_at
) VALUES (
    {sql_literal(title, tag + "_title")},
    {sql_literal(slug, tag + "_slug")},
    {sql_literal(excerpt, tag + "_excerpt")},
    {sql_literal(content, tag + "_content")},
    {sql_literal(thumbnail, tag + "_thumb")},
    {enrichment.category_id},
    'ALOO Editorial',
    'ALOO',
    {sql_literal(GOOGLE_SOURCE, tag + "_source")},
    'BLOG',
    'PUBLISHED',
    TIMESTAMP '{published_at}',
    {sql_literal(title, tag + "_seo_title")},
    {sql_literal(seo_description, tag + "_seo_description")},
    {sql_literal(tags, tag + "_keywords")},
    {sql_literal("/blog/" + slug, tag + "_canonical")},
    {sql_literal(tags, tag + "_tags")},
    TIMESTAMP '{published_at}',
    CURRENT_TIMESTAMP
)
ON CONFLICT (slug) DO UPDATE SET
    title = EXCLUDED.title,
    excerpt = EXCLUDED.excerpt,
    content = EXCLUDED.content,
    thumbnail_url = EXCLUDED.thumbnail_url,
    category_id = EXCLUDED.category_id,
    author = EXCLUDED.author,
    source = EXCLUDED.source,
    source_link = EXCLUDED.source_link,
    article_type = EXCLUDED.article_type,
    status = EXCLUDED.status,
    published_at = EXCLUDED.published_at,
    seo_title = EXCLUDED.seo_title,
    seo_description = EXCLUDED.seo_description,
    meta_keywords = EXCLUDED.meta_keywords,
    canonical_url = EXCLUDED.canonical_url,
    tags = EXCLUDED.tags,
    updated_at = CURRENT_TIMESTAMP;
DELETE FROM post_images WHERE post_id = (SELECT id FROM posts WHERE slug = {sql_literal(slug, tag + "_image_slug")});
""".strip()
        )

        for image_index, image_url in enumerate(image_urls):
            sql_parts.append(
                "INSERT INTO post_images (post_id, image_url, alt_text, image_type, sort_order, created_at) "
                f"SELECT id, {sql_literal(image_url, tag + f'_image_{image_index}')}, "
                f"{sql_literal(title, tag + f'_alt_{image_index}')}, 'GALLERY', {image_index}, CURRENT_TIMESTAMP "
                f"FROM posts WHERE slug = {sql_literal(slug, tag + f'_gallery_slug_{image_index}')};"
            )

        manifest_rows.append(
            f"{article_number}\t{title}\t{slug}\t{enrichment.category}\t{len(image_urls)}\t{len(content)}"
        )

    sql_parts.append("COMMIT;")
    OUTPUT_SQL.write_text("\n\n".join(sql_parts) + "\n", encoding="utf-8")
    (OUTPUT_SQL.parent / "blog-import-manifest.tsv").write_text(
        "\n".join(manifest_rows) + "\n",
        encoding="utf-8",
    )
    print(f"Prepared {len(start_indexes)} articles")
    print(f"SQL: {OUTPUT_SQL}")
    print(f"Images: {len(list(MEDIA_DIR.glob('aloo-blog-*')))}")
    print("\n".join(manifest_rows))


if __name__ == "__main__":
    main()
