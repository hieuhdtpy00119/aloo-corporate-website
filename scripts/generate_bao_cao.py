# -*- coding: utf-8 -*-
"""Generate project progress report as Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "Bao-cao-tien-do-ALOO-15072026.docx"


def set_run_font(run, name="Times New Roman", size=13, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_title(doc, text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=14 if level == 1 else 13, bold=True, color=(18, 35, 16))
    return h


def add_paragraph(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for run in hdr_cells[i].paragraphs[0].runs:
            set_run_font(run, bold=True, size=12)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for run in cells[i].paragraphs[0].runs:
                set_run_font(run, size=12)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    add_title(doc, "PHÒNG CÔNG NGHỆ THÔNG TIN", size=12)
    add_title(doc, "─" * 45, size=10)
    add_title(doc, "BÁO CÁO TIẾN ĐỘ DỰ ÁN", size=18)
    add_title(doc, "ALOO CORPORATE WEBSITE & FRANCHISE CMS", size=15)
    doc.add_paragraph()

    meta = [
        ("Tên dự án", "ALOO Corporate Website & Franchise CMS"),
        ("Đơn vị lập", "Phòng Công nghệ Thông tin"),
        ("Người lập", "Hiếu Huỳnh Đoàn Trung"),
        ("Ngày lập", date.today().strftime("%d/%m/%Y")),
        ("Phiên bản", "v1.0 – Báo cáo tiến độ off sớm"),
        ("Mục tiêu off", "15/07/2026 (sớm 2 tuần so với dự kiến trường)"),
        ("Mức độ", "NỘI BỘ"),
    ]
    add_table(doc, ["Hạng mục", "Nội dung"], meta)
    doc.add_paragraph()

    add_heading(doc, "I. TÓM TẮT ĐIỀU HÀNH", 1)
    add_paragraph(
        doc,
        "Báo cáo này tổng hợp tiến độ triển khai dự án ALOO Corporate Website & Franchise CMS "
        "nhằm phục vụ mục tiêu hoàn thành sớm vào ngày 15/07/2026 theo yêu cầu của Ban lãnh đạo. "
        "Dự án được phát triển theo mô hình 2 team: Frontend (Vue 3) và Backend (Spring Boot).",
    )
    add_paragraph(doc, "Đánh giá tổng thể:", bold=True)
    add_bullet(doc, "Phạm vi tính năng chính đã hoàn thiện, đủ điều kiện demo và bàn giao sớm.")
    add_bullet(doc, "Frontend build production thành công; unit test đạt 73/74 case.")
    add_bullet(doc, "Backend đã triển khai đầy đủ API cho public site, admin CMS và các module mới.")
    add_bullet(doc, "Cần tập trung 2 ngày cuối (13–15/07) cho QA tích hợp, migration DB và commit code còn lại.")

    add_heading(doc, "II. KIẾN TRÚC VÀ CÔNG NGHỆ", 1)
    add_table(
        doc,
        ["Tầng", "Công nghệ", "Vai trò", "Trạng thái"],
        [
            ("Frontend", "Vue 3 + Vite + Tailwind CSS", "Giao diện public & admin CMS", "Hoàn thiện"),
            ("State management", "Pinia", "Quản lý state ứng dụng", "Hoàn thiện"),
            ("Routing", "Vue Router", "Điều hướng trang", "Hoàn thiện"),
            ("Backend", "Spring Boot 3 (Java 21)", "API & nghiệp vụ", "Hoàn thiện"),
            ("Database", "SQL Server + JPA/Hibernate", "Lưu trữ dữ liệu", "Hoàn thiện"),
            ("Xác thực", "JWT + Google OAuth2", "Bảo mật đăng nhập", "Hoàn thiện"),
            ("Realtime", "WebSocket", "Live chat", "Hoàn thiện"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "III. TIẾN ĐỘ TEAM FRONTEND", 1)

    add_heading(doc, "3.1. Website công khai (Public)", 2)
    add_table(
        doc,
        ["Trang", "Route", "Chức năng chính", "Trạng thái"],
        [
            ("Trang chủ", "/", "Hero, block nổi bật, sản phẩm, cửa hàng, CTA", "Hoàn thành"),
            ("Giới thiệu", "/about", "Câu chuyện thương hiệu, timeline", "Hoàn thành"),
            ("Sản phẩm", "/products", "Danh sách, hero slider, menu poster", "Hoàn thành"),
            ("Chi tiết SP", "/products/:slug", "Thông tin SP, đánh giá khách hàng", "Hoàn thành"),
            ("Nhượng quyền", "/franchise", "Lợi ích, quy trình, chi phí, CTA", "Hoàn thành"),
            ("Cửa hàng", "/locations", "Tìm kiếm, lọc, bản đồ", "Hoàn thành"),
            ("Chi tiết CH", "/locations/:slug", "Thông tin chi nhánh", "Hoàn thành"),
            ("Blog", "/blog", "Danh sách bài viết, danh mục", "Hoàn thành"),
            ("Chi tiết blog", "/blog/:slug", "Nội dung, TOC, bài liên quan", "Hoàn thành"),
            ("Tư vấn", "/consultation", "Form đăng ký nhượng quyền", "Hoàn thành"),
            ("Liên hệ", "/contact", "Form liên hệ, thông tin hotline", "Hoàn thành"),
            ("Tài khoản", "/account", "Profile, đổi mật khẩu", "Hoàn thành"),
        ],
    )
    doc.add_paragraph()
    add_paragraph(doc, "Ghi chú: /process và /cost được redirect vào /franchise theo thiết kế mới.")

    add_heading(doc, "3.2. Hệ thống quản trị (Admin CMS)", 2)
    add_table(
        doc,
        ["Nhóm", "Module", "Chức năng", "Trạng thái"],
        [
            ("Dashboard", "Tổng quan", "Thống kê lead, bài viết, biểu đồ", "Hoàn thành"),
            ("Content", "Home Sections", "Quản lý block trang chủ", "Hoàn thành"),
            ("Content", "Products", "CRUD sản phẩm, menu poster", "Hoàn thành"),
            ("Content", "Franchise", "Nội dung trang nhượng quyền", "Hoàn thành"),
            ("Content", "Articles", "CRUD bài viết, SEO score", "Hoàn thành"),
            ("Content", "Categories", "Danh mục bài viết", "Hoàn thành"),
            ("Stores", "Locations", "CRUD cửa hàng, tiện ích", "Hoàn thành"),
            ("CRM", "Leads", "Quản lý đăng ký tư vấn", "Hoàn thành"),
            ("CRM", "Feedbacks", "Phản hồi khách hàng", "Hoàn thành"),
            ("CRM", "Product Reviews", "Duyệt đánh giá sản phẩm", "Hoàn thành"),
            ("CRM", "Live Chat", "Hỗ trợ chat realtime", "Hoàn thành"),
            ("System", "Accounts", "Quản lý tài khoản admin", "Hoàn thành"),
            ("System", "Audit Logs", "Nhật ký thao tác", "Hoàn thành"),
            ("System", "Profile", "Hồ sơ, bảo mật, OTP đổi MK", "Hoàn thành"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "3.3. Cải tiến UI/UX gần đây", 2)
    add_bullet(doc, "Thiết kế lại admin shell thống nhất (AdminShellFrame, nested layout, table panel).")
    add_bullet(doc, "Chuẩn hóa modal, breadcrumb, empty/loading/error state theo aloo-ui-ux-standard.")
    add_bullet(doc, "Đa ngôn ngữ VI/EN cho admin views.")
    add_bullet(doc, "Tích hợp ChatWidget trên public site.")
    add_bullet(doc, "Responsive mobile cho toàn bộ trang public.")

    add_heading(doc, "3.4. Kiểm thử Frontend", 2)
    add_table(
        doc,
        ["Hạng mục", "Kết quả", "Ghi chú"],
        [
            ("Unit test (Vitest)", "73/74 pass", "1 test AdminArticlesView cần cập nhật (skeleton loading)"),
            ("Production build", "Thành công", "Có warning CSS :deep và bundle >500KB"),
            ("E2E (Playwright)", "Đã có script", "public-user-flows + admin aloo-flows"),
            ("Test coverage admin", "4/17 màn", "Các module mới chưa có unit test"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "IV. TIẾN ĐỘ TEAM BACKEND", 1)

    add_heading(doc, "4.1. API đã triển khai", 2)
    add_bullet(doc, "Auth: đăng nhập admin/user, JWT, đổi mật khẩu, OTP, Google OAuth2.")
    add_bullet(doc, "Products, Posts, Categories: CRUD đầy đủ cho admin; public GET.")
    add_bullet(doc, "Stores/Locations: quản lý chi nhánh, featured stores, amenities JSON.")
    add_bullet(doc, "Franchise: nội dung landing, đăng ký tư vấn, contact messages.")
    add_bullet(doc, "Content: hero banners, menu posters, home sections, brand timeline.")
    add_bullet(doc, "CRM: feedbacks, product reviews (public submit + admin duyệt).")
    add_bullet(doc, "Live Chat: REST API + WebSocket cho phiên chat realtime.")
    add_bullet(doc, "System: account management, audit logs, upload ảnh, rate limiting.")

    add_heading(doc, "4.2. Database & Migration", 2)
    add_bullet(doc, "Schema chính: aloo_franchise_cms.sql + sample data.")
    add_bullet(doc, "Migration mới: product_reviews, auth_provider, live_chat, password_set_at, store_amenities.")
    add_bullet(doc, "RBAC theo admin scope: dashboard, content, stores, crm, system.")

    add_heading(doc, "4.3. Bảo mật", 2)
    add_bullet(doc, "JWT HS256, CORS cấu hình theo origin.")
    add_bullet(doc, "Rate limit: login, upload (memory backend cho local).")
    add_bullet(doc, "Phân quyền admin theo scope trên SecurityConfig.")
    add_bullet(doc, "Validation Jakarta trên DTO, không expose passwordHash.")

    add_heading(doc, "4.4. Kiểm thử Backend", 2)
    add_table(
        doc,
        ["Test suite", "Phạm vi", "Trạng thái"],
        [
            ("ApiSecurityIntegrationTest", "API smoke, JWT, CORS", "Có"),
            ("JwtServiceTest", "Token valid/expired", "Có"),
            ("ProductServiceTest", "CRUD, duplicate slug", "Có"),
            ("FranchiseRegistrationServiceTest", "Lead create/status", "Có"),
            ("ProductRepositoryDataJpaTest", "JPA repository", "Có"),
            ("AdminScopeCheckerTest", "RBAC scope", "Có"),
            ("PasswordChangePolicyTest", "OTP policy Google account", "Có"),
            ("Chat / Reviews / Audit", "Module mới", "Chưa có test riêng"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "V. TÍCH HỢP FE – BE", 1)
    add_table(
        doc,
        ["Luồng nghiệp vụ", "Trạng thái", "Ghi chú"],
        [
            ("Login admin → Dashboard", "Đã tích hợp", "Cần smoke test trước off"),
            ("CRUD sản phẩm qua admin UI", "Đã tích hợp", "E2E có script"),
            ("Form tư vấn → Lead admin CRM", "Đã tích hợp", "Field mapping camelCase OK"),
            ("Public products/locations/blog", "Đã tích hợp", "Có empty/error state"),
            ("Product reviews submit + duyệt", "Đã tích hợp", "Cần chạy migration DB"),
            ("Live chat WebSocket", "Đã tích hợp", "Cần test qua proxy deploy"),
            ("Google OAuth login", "Đã tích hợp", "Cần cấu hình client ID production"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "VI. COMMIT & THAY ĐỔI GẦN NHẤT", 1)
    add_paragraph(doc, "Các commit chính trên nhánh main:", bold=True)
    add_bullet(doc, "Add live chat, product reviews, password OTP, and admin shell redesign.")
    add_bullet(doc, "Deliver enterprise admin CMS with scoped RBAC, audit trail, and unified UX.")
    add_bullet(doc, "Add account management, contact message, brand timeline modules.")
    add_bullet(doc, "Refactor corporate site scope, CMS home sections, product pages.")

    add_paragraph(doc, "Code chưa commit (tính đến 13/07/2026):", bold=True)
    add_bullet(doc, "frontend/src/components/public/Footer.vue")
    add_bullet(doc, "frontend/src/views/public/HomeView.vue")
    add_bullet(doc, "frontend/src/views/admin/AdminLocationsView.vue")
    add_bullet(doc, "frontend/src/locales/admin-views-en.json")
    add_bullet(doc, "frontend/src/locales/admin-views-vi.json")

    add_heading(doc, "VII. HẠ TẦNG TRIỂN KHAI (ĐỀ XUẤT)", 1)
    add_paragraph(
        doc,
        "Theo tài liệu đề xuất hạ tầng (proposal v1.0, 08/06/2026), phương án được kiến nghị:",
    )
    add_table(
        doc,
        ["Hạng mục", "Đề xuất", "Chi phí/năm"],
        [
            ("Domain chính", "aloo.vn", "400.000 đ"),
            ("Domain phụ", "aloo.com (redirect)", "400.000 đ"),
            ("VPS", "PA2: 2 CPU – 4 GB RAM (AZDIGI)", "2.640.000 đ"),
            ("SSL", "Let's Encrypt", "Miễn phí"),
            ("CDN", "Cloudflare Free", "Miễn phí"),
            ("Tổng (có VAT)", "PA2", "3.784.000 đ"),
        ],
    )
    doc.add_paragraph()
    add_paragraph(
        doc,
        "Lưu ý: Tài liệu đề xuất ghi PostgreSQL, trong khi codebase hiện tại triển khai trên SQL Server. "
        "Cần thống nhất trước khi deploy production.",
        bold=False,
    )

    add_heading(doc, "VIII. VIỆC CẦN HOÀN THÀNH TRƯỚC OFF 15/07", 1)
    add_table(
        doc,
        ["STT", "Hạng mục", "Team", "Ưu tiên", "Deadline"],
        [
            ("1", "Commit 5 file FE đang pending", "FE", "Cao", "13/07"),
            ("2", "Fix 1 unit test AdminArticlesView", "FE", "Cao", "13/07"),
            ("3", "Smoke test 7 luồng tích hợp chính", "FE + BE", "Cao", "14/07"),
            ("4", "Chạy migration DB trên môi trường demo", "BE", "Cao", "14/07"),
            ("5", "Chạy mvn test full backend", "BE", "Cao", "14/07"),
            ("6", "Cập nhật README (bỏ ghi mock, port 8080)", "FE", "Trung bình", "14/07"),
            ("7", "Chuẩn bị tài khoản demo + sample data", "BE", "Trung bình", "14/07"),
            ("8", "Demo tích hợp full trước khi off", "Cả 2 team", "Cao", "15/07 sáng"),
        ],
    )
    doc.add_paragraph()

    add_heading(doc, "IX. KẾT LUẬN", 1)
    add_paragraph(
        doc,
        "Dự án ALOO Corporate Website & Franchise CMS đã hoàn thiện phạm vi tính năng chính cho cả "
        "website công khai và hệ thống quản trị CMS. Các module mới nhất gồm live chat, đánh giá sản phẩm, "
        "quản lý tài khoản, audit trail và đăng nhập Google đã được triển khai trên cả frontend và backend.",
    )
    add_paragraph(
        doc,
        "Với tiến độ hiện tại, dự án đủ điều kiện off sớm vào ngày 15/07/2026 nếu 2 team tập trung "
        "hoàn tất QA tích hợp, migration database và xử lý các hạng mục còn lại trong 2 ngày tới. "
        "Khuyến nghị không mở thêm tính năng mới, chỉ xử lý bug blocker trước khi bàn giao.",
    )
    doc.add_paragraph()

    add_paragraph(doc, "Người lập", bold=True)
    add_paragraph(doc, "Hiếu Huỳnh Đoàn Trung")
    add_paragraph(doc, "Phòng Công nghệ Thông tin")
    doc.add_paragraph()
    add_paragraph(doc, "Phê duyệt: ......................................")
    add_paragraph(doc, "Ban Lãnh đạo / Khách hàng (C)")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
